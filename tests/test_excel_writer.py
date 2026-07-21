"""Tests for the Excel writer, focusing on the override round-trip."""

import tempfile
from pathlib import Path

from openpyxl import Workbook

from src.db import get_connection, init_db
from src.excel_writer import read_overrides_from_excel, TXN_LOG_HEADERS, TXN_LOG_START_ROW


def _create_test_db(tmpdir):
    """Create a test DB with schema and a sample transaction."""
    db_path = Path(tmpdir) / "test.db"
    init_db(db_path)
    conn = get_connection(db_path)

    conn.execute(
        "INSERT INTO accounts (id, institution, name) VALUES ('A1', 'Test Bank', 'Checking')"
    )
    conn.execute(
        """INSERT INTO transactions
           (id, account_id, posted_at, amount, description, pending, tier1, tier2, vendor, auto_categorized, month)
           VALUES ('TXN-001', 'A1', '2026-03-15', -42.50, 'SOME STORE #123', 0,
                   'Non-Essential', 'Shopping', 'Some Store', 1, '2026-03')"""
    )
    conn.commit()
    return db_path, conn


def _create_test_workbook(tmpdir, txn_id, tier1, tier2, reason, override_flag="Y"):
    """Create a minimal workbook with a Transaction Log sheet containing one overridden row."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Transaction Log"

    for col_idx, header in enumerate(TXN_LOG_HEADERS, start=1):
        ws.cell(row=TXN_LOG_START_ROW, column=col_idx, value=header)

    data_row = TXN_LOG_START_ROW + 1
    ws.cell(row=data_row, column=1, value="2026-03-15")
    ws.cell(row=data_row, column=2, value="Some Store")
    ws.cell(row=data_row, column=3, value=-42.50)
    ws.cell(row=data_row, column=4, value="Test Bank - Checking")
    ws.cell(row=data_row, column=5, value=tier1)
    ws.cell(row=data_row, column=6, value=tier2)
    ws.cell(row=data_row, column=7, value="2026-03")
    ws.cell(row=data_row, column=8, value=reason)
    ws.cell(row=data_row, column=9, value=override_flag)
    ws.cell(row=data_row, column=10, value=txn_id)

    excel_path = Path(tmpdir) / "test_dashboard.xlsx"
    wb.save(str(excel_path))
    wb.close()
    return excel_path


def test_override_round_trip_preserves_vendor():
    """Override from Excel should set reason, NOT corrupt the vendor column."""
    with tempfile.TemporaryDirectory() as tmpdir:
        db_path, conn = _create_test_db(tmpdir)

        excel_path = _create_test_workbook(
            tmpdir,
            txn_id="TXN-001",
            tier1="Essential",
            tier2="Groceries",
            reason="Actually a grocery run",
        )

        count = read_overrides_from_excel(conn, excel_path)
        assert count == 1

        txn = conn.execute(
            "SELECT vendor, tier1, tier2 FROM transactions WHERE id = 'TXN-001'"
        ).fetchone()
        assert txn["tier1"] == "Essential"
        assert txn["tier2"] == "Groceries"
        assert txn["vendor"] is None or txn["vendor"] != "Actually a grocery run"

        override = conn.execute(
            "SELECT vendor, reason FROM overrides WHERE transaction_id = 'TXN-001'"
        ).fetchone()
        assert override["reason"] == "Actually a grocery run"
        assert override["vendor"] is None

        conn.close()


def test_override_without_reason_uses_default():
    """Override with no Notes column value should use 'Excel override' as reason."""
    with tempfile.TemporaryDirectory() as tmpdir:
        db_path, conn = _create_test_db(tmpdir)

        excel_path = _create_test_workbook(
            tmpdir,
            txn_id="TXN-001",
            tier1="Essential",
            tier2="Medical",
            reason=None,
        )

        count = read_overrides_from_excel(conn, excel_path)
        assert count == 1

        override = conn.execute(
            "SELECT vendor, reason FROM overrides WHERE transaction_id = 'TXN-001'"
        ).fetchone()
        assert override["reason"] == "Excel override"
        assert override["vendor"] is None

        conn.close()


def test_override_flag_not_set_skips_row():
    """Rows without Override=Y should not be processed."""
    with tempfile.TemporaryDirectory() as tmpdir:
        db_path, conn = _create_test_db(tmpdir)

        excel_path = _create_test_workbook(
            tmpdir,
            txn_id="TXN-001",
            tier1="Essential",
            tier2="Groceries",
            reason="Should not apply",
            override_flag="",
        )

        count = read_overrides_from_excel(conn, excel_path)
        assert count == 0

        txn = conn.execute(
            "SELECT tier1, tier2, vendor FROM transactions WHERE id = 'TXN-001'"
        ).fetchone()
        assert txn["tier1"] == "Non-Essential"
        assert txn["tier2"] == "Shopping"
        assert txn["vendor"] == "Some Store"

        conn.close()


# ─── Formula generation tests ───

from src.excel_writer import (
    _build_formula,
    _sumifs_spend,
    _sumifs_alloc,
    _sumifs_income,
    _sumifs_income_other,
    _needs_row_insertion,
)


def test_sumifs_spend_single_tier2():
    """Spend formula for a single tier2 category."""
    result = _sumifs_spend(["Groceries"], "C")
    assert result.startswith("=-")
    assert '"Groceries"' in result
    assert "C$3" in result


def test_sumifs_spend_multiple_tier2s():
    """Spend formula joins multiple tier2 categories with subtraction."""
    result = _sumifs_spend(["Dining Out", "Food Delivery"], "D")
    assert result.count("SUMIFS") == 2
    assert '"Dining Out"' in result
    assert '"Food Delivery"' in result


def test_sumifs_spend_empty_tier2_returns_zero():
    """Empty tier2 list should produce =0."""
    result = _sumifs_spend([], "C")
    assert result == "=0"


def test_sumifs_alloc_formula():
    """Alloc formula references tier1=Allocation."""
    result = _sumifs_alloc(["Savings"], "E")
    assert '"Allocation"' in result
    assert '"Savings"' in result


def test_sumifs_income_formula():
    """Income formula references tier1=Income, tier2=Payroll."""
    result = _sumifs_income("C")
    assert '"Income"' in result
    assert '"Payroll"' in result


def test_sumifs_income_other_formula():
    """Income-other subtracts Payroll from total Income."""
    result = _sumifs_income_other("F")
    assert result.count("SUMIFS") == 2
    assert '"Income"' in result
    assert '"Payroll"' in result


def test_build_formula_sum_type():
    """Sum type produces SUM(col+first:col+last)."""
    result = _build_formula({"type": "sum", "first": 12, "last": 17}, "C")
    assert result == "=SUM(C12:C17)"


def test_build_formula_ref_type():
    """Ref type produces =col+row."""
    result = _build_formula({"type": "ref", "row": 8}, "D")
    assert result == "=D8"


def test_build_formula_template():
    """Formula type replaces {col} placeholder."""
    result = _build_formula(
        {"type": "formula", "template": "={col}71-{col}72"}, "G"
    )
    assert result == "=G71-G72"


# ─── Row insertion detection ───

def test_needs_row_insertion_false_when_rideshare():
    """When B45 is already Rideshare, no insertion needed."""
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.cell(row=45, column=2, value="Rideshare")
    assert _needs_row_insertion(ws) is False
    wb.close()


def test_needs_row_insertion_true_when_different():
    """When B45 is something else, insertion IS needed."""
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.cell(row=45, column=2, value="Something Else")
    assert _needs_row_insertion(ws) is True
    wb.close()


def test_needs_row_insertion_true_when_empty():
    """When B45 is empty, insertion IS needed."""
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    assert _needs_row_insertion(ws) is True
    wb.close()


if __name__ == "__main__":
    import pytest
    pytest.main([__file__, "-v"])
